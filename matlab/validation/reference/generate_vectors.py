"""
Generates the P.619-6 validation vectors in ../vectors/ from the
independent reference implementation p619_ref.py and, where available,
from the official ITU-R Study Group 3 validation examples
(CG-3M3J-13-ValEx-Rev8.1.1.xlsx).

Usage:  python3 generate_vectors.py [path/to/CG-3M3J-13-ValEx-Rev8.1.1.xlsx]

Requires: numpy, scipy, itur (ITU-Rpy), openpyxl
"""

import csv
import os
import sys

import numpy as np
import openpyxl
from scipy.optimize import brentq

import p619_ref as R

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, '..', 'vectors')
XLSX = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
    HERE, '..', '..', 'matlab', 'validation_examples', 'CG-3M3J-13-ValEx-Rev8.1.1.xlsx')

os.makedirs(OUT, exist_ok=True)


def write(name, header, rows):
    with open(os.path.join(OUT, name), 'w', newline='') as fh:
        w = csv.writer(fh, lineterminator='\n')
        w.writerow(header)
        for k, r in enumerate(rows, start=1):
            w.writerow([k] + [x if isinstance(x, str) else '%.15g' % x for x in r])
    print('%-32s %4d rows' % (name, len(rows)))


def wrap180(x):
    return (x + 180.0) % 360.0 - 180.0


wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)

# 01 Free-space basic transmission loss, eq. (1) -----------------------------
rows = [(f, d, R.lbfs(f, d), 'eq. (1)') for f, d in
        [(0.1, 1.0), (1.0, 500.0), (6.425, 38134.0), (14.25, 36000.0), (20.0, 550.0),
         (28.0, 2000.0), (50.0, 42164.0), (100.0, 100.0)]]
write('01_free_space.csv', ['id', 'f_GHz', 'd_km', 'Lbfs_dB', 'source'], rows)

# 02 Depolarisation from XPD, eqs. (2a), (2b), (6) ----------------------------
rows = []
for xpd in [-10.0, 0.0, 3.0, 10.0, 15.0, 20.0, 30.0, 40.0]:
    ax, ac = R.xpd_attenuation(xpd)
    rows.append((xpd, ax, ac, R.hydrometeor_depol_attenuation(xpd), 'eqs. (2a), (2b), (6)'))
write('02_depolarisation.csv', ['id', 'XPD_dB', 'Ax_dB', 'Ac_dB', 'Axq_dB', 'source'], rows)

# 03 Faraday rotation, eqs. (3a), (3b), (4) ------------------------------------
rows = []
for f, bav, nt in [(0.1, 50e-6, 1e16), (0.4, 50e-6, 1e17), (1.0, 50e-6, 1e17), (1.5, 30e-6, 1e18),
                   (2.0, 50e-6, 1e18), (4.0, 50e-6, 1e18), (6.0, 60e-6, 1e19), (10.0, 50e-6, 1e19)]:
    th = R.faraday_rotation(f, bav, nt)
    if not 0.0 < th < np.pi / 2:
        continue
    axf, acf = R.faraday_attenuation_literal(th)
    rows.append((f, bav, nt, th, axf, acf, 'eqs. (3a), (3b), (4) as printed'))
write('03_faraday.csv', ['id', 'f_GHz', 'Bav_T', 'NT_el_per_m2', 'thetaF_rad', 'AxF_dB', 'AcF_dB', 'source'], rows)

# 04 Beam spreading, P.834 §5 (P.619-6 §2.4.2) -------------------------------
rows = [(t, h, R.beam_spreading_loss(t, h), 'P.834 §5') for h in (0.0, 0.5, 1.0, 2.0, 3.0, 5.0)
        for t in (0.0, 0.5, 1.0, 2.0, 5.0, 10.0)]
write('04_beam_spreading.csv', ['id', 'theta0_deg', 'h_km', 'Abs_dB', 'source'], rows)

# 05 Attachment A geometry --------------------------------------------------
cases = [  # Hs, Ht, lat_s, lon_s, lat_t, lon_t
    (35786.0, 0.55, 0.0, 0.0, 46.95, 7.44),     # Bern -> GSO 0E
    (35786.0, 0.02, 0.0, 64.0, 51.5072, -0.1276),  # London -> GSO 64E
    (35786.0, 0.1, 0.0, -20.0, 69.65, 18.96),   # Tromso, low elevation
    (35786.0, 0.1, 0.0, -60.0, 69.65, 18.96),   # Tromso, below horizon
    (35786.0, 0.0, 0.0, -100.0, -22.9, -43.23),  # Rio de Janeiro (southern hemisphere)
    (35786.0, 0.0, 0.0, 170.0, 0.0, -175.0),    # across the antimeridian
    (550.0, 0.55, 48.0, 10.0, 46.95, 7.44),     # LEO, high elevation
    (550.0, 0.55, 52.0, 18.0, 46.95, 7.44),     # LEO, mid elevation
    (1200.0, 2.0, 60.0, -100.0, 45.0, -90.0),   # LEO north of station
    (20.0, 0.55, 47.5, 8.0, 46.95, 7.44),       # HAPS within the atmosphere
    (35786.0, 0.0, 0.0, 0.0, 0.0, 0.0),         # sub-satellite point (zenith)
]
rows = []
for hs, ht, la_s, lo_s, la_t, lo_t in cases:
    delta = wrap180(lo_s - lo_t)
    d, th0, psi = R.straight_line_geometry(hs, ht, la_s, la_t, delta)
    d2, th2, psi2 = R.look_angles_ecef(hs, ht, la_s, lo_s, la_t, lo_t)
    assert abs(d - d2) < 1e-8 and abs(th0 - th2) < 1e-9
    if th0 < 89.999:
        assert abs(wrap180(psi - psi2)) < 1e-8
    else:
        psi = float('nan')  # bearing indeterminate for a vertical path
    rows.append((hs, ht, la_s, lo_s, la_t, lo_t, delta, d, th0, psi, 'eqs. (18)-(24); checked vs ECEF/ENU'))
write('05_geometry_attA.csv', ['id', 'Hs_km', 'Ht_km', 'lat_s_deg', 'lon_s_deg', 'lat_t_deg', 'lon_t_deg',
                               'delta_deg', 'Dts_km', 'theta0_deg', 'psi_deg', 'source'], rows)

# 06/07 Attachment B ---------------------------------------------------------
rows = [(ht, t0, R.fs_to_apparent(t0, ht), 'eqs. (25)-(26)') for ht in (0.0, 0.5, 1.0, 2.0, 3.0)
        for t0 in (-1.0, 0.0, 1.0, 2.0, 5.0, 10.0)]
write('06_attB_fs_to_apparent.csv', ['id', 'Ht_km', 'theta0_deg', 'theta_deg', 'source'], rows)
rows = [(ht, t, R.apparent_to_fs(t, ht), 'eqs. (27)-(28)') for ht in (0.0, 0.5, 1.0, 2.0, 3.0)
        for t in (0.0, 1.0, 2.0, 5.0, 10.0)]
write('07_attB_apparent_to_fs.csv', ['id', 'Ht_km', 'theta_deg', 'theta0_deg', 'source'], rows)

# 08 Specific attenuation, P.676 Annex 1 eq. (1) -----------------------------
rows = []
ws = list(wb['P.676-13 SpAtt'].iter_rows(values_only=True))
freqs = [12.0, 20.0, 60.0, 90.0, 130.0]
go = [v for v in ws[27] if isinstance(v, float)][:5]
gw = [v for v in ws[27] if isinstance(v, float)][5:10]
for f, a, b in zip(freqs, go, gw):
    rows.append((f, 1013.25, 7.5, 288.15, a, b, 'ITU ValEx P.676-13 SpAtt'))
for f, pd, rho, t in [(1.0, 1013.25, 7.5, 288.15), (10.0, 1013.25, 7.5, 288.15), (22.235, 1013.25, 7.5, 288.15),
                      (50.0, 1013.25, 7.5, 288.15), (57.0, 900.0, 5.0, 280.0), (118.75, 1013.25, 7.5, 288.15),
                      (183.31, 1013.25, 7.5, 288.15), (20.0, 500.0, 0.5, 250.0), (30.0, 100.0, 0.01, 220.0),
                      (60.0, 10.0, 1e-5, 230.0), (300.0, 1013.25, 12.0, 300.0)]:
    a, b = R.specific_attenuation(f, pd, rho, t)
    rows.append((f, pd, rho, t, a, b, 'ITU-Rpy itu676 (P.676-12 lines)'))
write('08_specific_attenuation.csv', ['id', 'f_GHz', 'pdry_hPa', 'rho_gm3', 'T_K', 'gamma_o_dBkm', 'gamma_w_dBkm', 'source'], rows)

# 09 MAGRA profiles + refractive index --------------------------------------
rows = []
for h in (0.0, 1.0, 2.0, 5.0, 10.0, 11.0, 15.0, 20.0, 25.0, 32.0, 40.0, 47.0, 50.0, 60.0, 71.0, 80.0, 85.5, 86.5, 90.0, 95.0, 100.0):
    t, p, rho, e = R.magra(h)
    rows.append((h, t, p, rho, e, R.refractive_index(t, p, e), 'P.835-6 §1.1 (MAGRA), P.453 eqs. (1),(6)'))
write('09_atmosphere_magra.csv', ['id', 'h_km', 'T_K', 'P_hPa', 'rho_gm3', 'e_hPa', 'n', 'source'], rows)

# 10 Attachment C.5 Case 1: ascending ray -------------------------------------
rows = []
for f, he, hs, phi, ref, src in [
        (28.0, 0.0, 100.0, 30.0, 0.47081173472870474, 'ITU ValEx P.676-13 A_Gas_A1_2.2.1a'),
        (28.0, 1.3, 8.0, 30.0, 0.24376211236218553, 'ITU ValEx P.676-13 A_Gas_A1_2.2.1b ex. 1'),
        (28.0, 1.3, 100.0, 30.0, 0.2774411060456813, 'ITU ValEx P.676-13 A_Gas_A1_2.2.1b ex. 2')]:
    rows.append((f, he, hs, phi, ref, src))
for f in (1.0, 10.0, 20.0, 28.0, 40.0, 50.0, 60.0, 90.0):
    for he, phi in [(0.0, 90.0), (0.0, 30.0), (0.0, 10.0), (0.0, 5.0), (0.0, 1.0), (0.0, 0.0),
                    (1.0, 20.0), (3.0, 2.0)]:
        rows.append((f, he, 35786.0, phi, R.ag_ascending(f, he, 35786.0, phi), 'reference (Hs above 100 km)'))
for f, he, hs, phi in [(20.0, 0.0, 20.0, 45.0), (40.0, 0.5, 12.0, 10.0), (60.0, 10.0, 100.0, 5.0)]:
    rows.append((f, he, hs, phi, R.ag_ascending(f, he, hs, phi), 'reference (Hs within atmosphere)'))
write('10_ag_ascending.csv', ['id', 'f_GHz', 'He_km', 'Hs_km', 'phi_e_deg', 'Ag_dB', 'source'], rows)

# 11 Attachment C.4: descending ray -----------------------------------------
rows = []
for f, he, hs, th in [(20.0, 0.0, 100.0, 30.0), (20.0, 0.0, 35786.0, 30.0), (28.0, 0.0, 35786.0, 10.0),
                      (40.0, 0.5, 35786.0, 5.0), (60.0, 1.0, 550.0, 20.0), (10.0, 0.0, 35786.0, 1.0),
                      (14.25, 0.2, 20.0, 15.0)]:
    phs = R.space_launch_angle(he, hs, th)
    rows.append((f, he, hs, th, phs, R.ag_descending(f, he, hs, phs), 'reference'))
write('11_ag_descending.csv', ['id', 'f_GHz', 'He_km', 'Hs_km', 'phi_e_arrival_deg', 'phi_s_deg', 'Ag_dB', 'source'], rows)

# 12/13 Attachment C.5 Case 2: negative elevation ------------------------------
rows_h, rows_a = [], []
for he, phi in [(30.0, -2.0), (10.0, -1.0), (10.0, -3.0), (5.0, -0.5), (5.0, -2.0), (2.0, -0.2), (2.0, -1.0), (1.0, -0.3)]:
    try:
        hmin = R.hmin_eq45(he, phi)
    except ValueError:
        continue  # ray intersects the Earth (no root with Hmin >= 0)
    rows_h.append((he, phi, hmin, 'eq. (45), Brent root'))
    for f in (20.0, 60.0):
        _, ag = R.ag_negative_elevation(f, he, 35786.0, phi)
        rows_a.append((f, he, 35786.0, phi, hmin, ag, 'eq. (44), reference'))
write('12_hmin.csv', ['id', 'He_km', 'phi_e_deg', 'Hmin_km', 'source'], rows_h)
write('13_ag_negative_elevation.csv', ['id', 'f_GHz', 'He_km', 'Hs_km', 'phi_e_deg', 'Hmin_km', 'Ag_dB', 'source'], rows_a)

# 14 Attachment D tropospheric scintillation ------------------------------------
rows = []
for r in wb['P.618-14 A_Scint'].iter_rows(min_row=21, values_only=True):
    v = [x for x in r if x is not None]
    if len(v) < 16 or not all(isinstance(x, (int, float)) for x in v[:16]):
        continue
    f, th, dia, eta, p618, nwet, sig, as618 = v[4], v[5], v[6], v[7], v[8], v[9], v[14], v[15]
    ga = 20.0 * np.log10(np.sqrt(eta) * dia * np.pi * f / 0.3)  # Deff of eq. (48) = sqrt(eta)*D
    rows.append((f, th, nwet, ga, 100.0 - p618, sig, as618, 'ITU ValEx P.618-14 A_Scint (fade side)'))
for f, th, nwet, ga in [(12.0, 10.0, 40.0, 35.0), (20.0, 30.0, 60.0, 45.0), (30.0, 60.0, 100.0, 50.0), (50.0, 5.0, 80.0, 30.0)]:
    for p in (0.001, 0.01, 0.1, 1.0, 10.0, 50.0, 90.0, 99.0, 99.9, 99.999):
        rows.append((f, th, nwet, ga, p, R.scintillation_sigma(f, th, nwet, ga),
                     R.scintillation_ast(f, p, th, nwet, ga), 'reference'))
write('14_scintillation_attD.csv', ['id', 'f_GHz', 'theta_deg', 'Nwet', 'Ga_dBi', 'p_pct', 'sigma_dB', 'Ast_dB', 'source'], rows)

# 15 Attachment E ray height -----------------------------------------------
rows = []
for ht in (0.0, 0.5, 2.0):
    for th in (0.0, 0.5, 1.0, 2.0, 5.0):
        for dc in (1, 10, 50, 100, 200):
            hr = R.ray_height(ht, th, dc)
            if hr < 10.0:
                rows.append((ht, th, dc, hr, 'eqs. (51)-(58)'))
    for th in (10.0, 30.0):
        for dc in (1, 5, 10, 20):
            hr = R.ray_height(ht, th, dc)
            if hr < 10.0:
                rows.append((ht, th, dc, hr, 'eq. (59)'))
write('15_ray_height_attE.csv', ['id', 'Ht_km', 'theta_deg', 'Dc_km', 'Hr_km', 'source'], rows)

# 16 §2.6 diffraction/ducting ----------------------------------------------
rows = []
# Cases chosen so that 0 < Ld + A(p) + Ads < Ld (formula fully exercised),
# plus three cases that exercise the clamps and the p >= beta branch.
for f, p, lat, thh, dh, ld, kind in [
        (2.0, 1.0, 47.0, 0.0, 5.0, 20.0, 'unclamped, Ads = 0'),
        (2.0, 3.0, 47.0, 0.0, 5.0, 20.0, 'unclamped, Ads = 0'),
        (10.0, 0.5, 20.0, 0.0, 2.0, 25.0, 'unclamped, Ads = 0'),
        (20.0, 2.0, 60.0, 0.0, 10.0, 15.0, 'unclamped, Ads = 0'),
        (40.0, 0.1, 80.0, 0.0, 10.0, 30.0, 'unclamped, beta = 4.17 %'),
        (6.0, 0.01, 35.0, 1.0, 5.0, 25.0, 'unclamped, Ads > 0'),
        (12.0, 0.1, 47.0, 2.0, 10.0, 20.0, 'unclamped, Ads > 0'),
        (28.0, 1.0, 10.0, 1.0, 8.0, 15.0, 'unclamped, Ads > 0'),
        (1.0, 0.001, 70.0, 3.0, 1.0, 30.0, 'unclamped, Ads > 0, |lat| = 70'),
        (2.0, 0.001, 47.0, 10.0, 5.0, 15.0, 'clamped at Ld'),
        (6.0, 0.001, 35.0, 0.0, 1.0, 5.0, 'clamped at 0'),
        (20.0, 10.0, 47.0, 20.0, 10.0, 25.0, 'p >= beta')]:
    v = R.ldtb(f, p, lat, thh, dh, ld)
    if kind.startswith('unclamped'):
        assert 0.0 < v < ld, (f, p, lat, thh, dh, ld, v)
    rows.append((f, p, lat, thh, dh, ld, v, 'eqs. (13)-(13e); ' + kind))
write('16_ducting_ldtb.csv', ['id', 'f_GHz', 'p_pct', 'lat_deg', 'theta_hoz_mrad', 'd_hoz_km', 'Ld_dB', 'Ldtb_dB', 'source'], rows)

# 18 Attachment F: precipitation-scatter test -------------------------------
# Typical Earth-station / terrestrial geometries; r at the -3 dB beam edges.
# Pint is set 5 dB either side of the decision threshold (Ptxs + 20 dB) so the
# decision is robust to round-off. Pairs with r1 = r2 and cyl = 1 / 2 check
# the symmetry that motivates the (60b) correction (identical Ptxs expected).
cases_f = [  # f, r1, r2, Ptx, G, dtx, drx, gamma_g, cyl, theta, theta_path, R, tau
    (5.0, 150.0, 400.0, 10.0, 45.0, 8.0, 15.0, 0.01, 1, 30.0, 5.0, 30.0, 0.0),
    (10.0, 150.0, 400.0, 10.0, 45.0, 8.0, 15.0, 0.012, 2, 30.0, 5.0, 30.0, 90.0),
    (14.25, 200.0, 200.0, 13.0, 50.0, 5.0, 12.0, 0.02, 1, 45.0, 10.0, 50.0, 45.0),
    (14.25, 200.0, 200.0, 13.0, 50.0, 5.0, 12.0, 0.02, 2, 45.0, 10.0, 50.0, 45.0),
    (20.0, 100.0, 300.0, 5.0, 55.0, 3.0, 20.0, 0.1, 1, 60.0, 20.0, 25.0, 90.0),
    (20.0, 100.0, 300.0, 5.0, 55.0, 3.0, 20.0, 0.1, 2, 60.0, 20.0, 25.0, 90.0),
    (28.0, 80.0, 80.0, 0.0, 60.0, 4.0, 6.0, 0.15, 1, 20.0, 30.0, 80.0, 0.0),
    (28.0, 80.0, 80.0, 0.0, 60.0, 4.0, 6.0, 0.15, 2, 20.0, 30.0, 80.0, 0.0),
    (30.0, 50.0, 500.0, -3.0, 52.0, 10.0, 2.0, 0.18, 2, 75.0, 2.0, 40.0, 45.0),
    (40.0, 120.0, 250.0, 7.0, 48.0, 2.0, 7.0, 0.25, 1, 90.0, 45.0, 60.0, 90.0),
    (40.0, 120.0, 250.0, 7.0, 48.0, 2.0, 7.0, 0.25, 2, 10.0, 45.0, 60.0, 0.0),
    (50.0, 300.0, 300.0, 3.0, 50.0, 1.0, 1.5, 0.5, 2, 35.0, 60.0, 100.0, 45.0),
]
rows = []
for k, (f, r1, r2, ptx, g, dtx, drx, gg, cyl, th, thp, rr, tau) in enumerate(cases_f):
    ptxs = R.precipitation_scatter_test(0.0, f, r1, r2, ptx, g, dtx, drx, gg, cyl, th, thp, rr, tau)['Ptxs_dBW']
    pint = round(ptxs + (15.0 if k % 2 == 0 else 25.0), 1)
    o = R.precipitation_scatter_test(pint, f, r1, r2, ptx, g, dtx, drx, gg, cyl, th, thp, rr, tau)
    rows.append((pint, f, r1, r2, ptx, g, dtx, drx, gg, cyl, th, thp, rr, tau,
                 o['L1_m'], o['L2_m'], o['S_dBWm2'], o['Pin_dBW'], o['gamma_r_dBkm'], o['Pout_dBW'],
                 o['Pscat_dBW'], o['Peirps_dBW'], o['Fnis_dB'], o['Ptxs_dBW'], o['full_calc'],
                 'eqs. (60a)-(67), (60b) with 2*r1/sin(theta)'))
# equal-beam symmetry pairs: rows 3/4 and 7/8
assert abs(rows[2][23] - rows[3][23]) < 1e-9 and abs(rows[6][23] - rows[7][23]) < 1e-9
write('18_precip_scatter_attF.csv',
      ['id', 'Pint_dBW', 'f_GHz', 'r1_m', 'r2_m', 'Ptx_dBW', 'G_dBi', 'dtx_km', 'drx_km', 'gamma_g_dBkm',
       'cyl', 'theta_deg', 'theta_path_deg', 'R_mmh', 'tau_deg',
       'L1_m', 'L2_m', 'S_dBWm2', 'Pin_dBW', 'gamma_r_dBkm', 'Pout_dBW', 'Pscat_dBW', 'Peirps_dBW',
       'Fnis_dB', 'Ptxs_dBW', 'full_calc', 'source'], rows)

# 19 §2.6 Fresnel parameters and knife-edge loss --------------------------------
rows = []
for h, d1, f, note in [(60.3, 24.0, 30.0, 'P.619-6 Fig. 10 example (text quotes nu ~ 4.9)'),
                       (0.0, 10.0, 20.0, 'grazing'), (-5.0, 5.0, 12.0, 'clearance'), (-20.0, 2.0, 30.0, 'clearance, nu < -0.78'),
                       (-30.0, 1.0, 6.0, 'outside first Fresnel zone'), (10.0, 0.5, 2.0, ''), (25.0, 15.0, 1.0, ''),
                       (150.0, 40.0, 50.0, 'deep shadow'), (-8.0, 12.0, 28.0, ''), (3.0, 30.0, 100.0, '')]:
    nu, r1 = R.fresnel_parameters(h, d1, f)
    rows.append((h, d1, f, nu, r1, R.knife_edge_loss(nu), 'eqs. (12a), (12b); P.526-16 eq. (31)' + ('; ' + note if note else '')))
write('19_fresnel_knife_edge.csv', ['id', 'h_m', 'd1_km', 'f_GHz', 'nu', 'R1_m', 'J_dB', 'source'], rows)

# 20 §2.6 finite-width screen (P.526-16 §5.1) ----------------------------------
rows = []
for nt, nl, nr in [(2.0, 50.0, 50.0), (2.0, 2.0, 2.0), (0.0, 5.0, 5.0), (5.0, 1.0, 8.0), (10.0, 3.0, 3.0),
                   (1.0, 0.5, 20.0), (3.0, -0.5, 4.0), (-0.9, 6.0, 6.0)]:
    jmin, jav = R.finite_width_screen_loss(nt, nl, nr)
    rows.append((nt, nl, nr, jmin, jav, 'P.526-16 eqs. (71), (72)'))
write('20_finite_width_screen.csv', ['id', 'nu_top', 'nu_left', 'nu_right', 'Jmin_dB', 'Jav_dB', 'source'], rows)

# 21 §2.6 terrain obstruction and ducting, eq. (13) -----------------------------
# Synthetic ridge profile d_i = i*step, h_i = h0 + A exp(-((d_i - dc)/w)^2) (m asl)
cases_t = [  # f, p, lat, Ht, theta, h0, A, dc, w, dmax, step, regime
    (20.0, 0.01, 47.0, 0.35, 1.0, 300.0, 300.0, 20.0, 3.0, 60.0, 0.1, 'unobstructed'),
    (6.0, 0.001, 30.0, 0.02, 0.2, 0.0, 150.0, 45.0, 4.0, 90.0, 0.25, 'unobstructed'),
    (20.0, 0.01, 47.0, 0.35, 1.0, 300.0, 600.0, 20.0, 3.0, 60.0, 0.1, 'obstructed, high horizon: Ldtb = Ld'),
    (14.25, 0.1, 47.0, 0.5, 2.0, 400.0, 700.0, 12.0, 2.0, 40.0, 0.1, 'obstructed, high horizon: Ldtb = Ld'),
    (30.0, 0.001, 47.0, 0.05, -0.1, 0.0, 120.0, 24.0, 2.0, 60.0, 0.1, 'partial ducting reduction'),
    (30.0, 0.01, 47.0, 0.05, -0.1, 0.0, 120.0, 24.0, 2.0, 60.0, 0.1, 'partial ducting reduction'),
    (30.0, 0.001, 47.0, 0.05, -0.1, 0.0, 75.0, 24.0, 2.0, 60.0, 0.1, 'partial, theta_hoz < 0'),
    (30.0, 0.001, 47.0, 0.05, -0.1, 0.0, 85.0, 24.0, 2.0, 60.0, 0.1, 'partial'),
    (10.0, 0.01, 20.0, 0.02, 0.0, 0.0, 120.0, 35.0, 3.0, 80.0, 0.2, 'clamped at 0'),
    (1.0, 0.01, 75.0, 0.05, 0.2, 0.0, 300.0, 40.0, 3.0, 80.0, 0.2, 'clamped at 0, |lat| > 70'),
    (28.0, 20.0, 47.0, 0.05, -0.1, 0.0, 120.0, 24.0, 2.0, 60.0, 0.1, 'p >= beta: Ldtb = Ld'),
    (12.0, 0.01, 47.0, 1.0, 8.0, 900.0, 1500.0, 8.0, 1.5, 30.0, 0.05, 'theta > 5 deg: eq. (59)'),
]
AE = 6371.0 * 4.0 / 3.0
rows = []
for f, p, lat, ht, th, h0, a, dc, w, dmax, st, reg in cases_t:
    d, h = R.ridge_profile(h0, a, dc, w, dmax, st)
    o = R.diffraction_ducting(f, p, lat, ht, th, d, h, AE)
    rows.append((f, p, lat, ht, th, AE, h0, a, dc, w, dmax, st, o['nu_max'], o['d_max'], o['h_max'], o['Ld'],
                 o['th_hoz'], o['d_hoz'], o['Ldtb'], '§2.6 eq. (13), single knife-edge method; ' + reg))
write('21_terrain_ldtb.csv', ['id', 'f_GHz', 'p_pct', 'lat_deg', 'Ht_km', 'theta_deg', 'ae_km', 'h0_m', 'A_m', 'dc_km',
                              'w_km', 'dmax_km', 'step_km', 'nu_max', 'd_max_km', 'h_max_m', 'Ld_dB', 'theta_hoz_mrad',
                              'd_hoz_km', 'Ldtb_dB', 'source'], rows)

# 22 §2.6 terrain loss with the P.526-16 §4.6 slant-path method ------------------
# Profile with up to two ridges: h_i = h0 + sum_k A_k exp(-((d_i - dc_k)/w_k)^2), d_i = i*step
cases_s = [  # f, p, lat, Ht, theta, h0, (A1, dc1, w1), (A2, dc2, w2), dmax, step, regime
    (20.0, 0.01, 47.0, 0.35, 1.0, 300.0, (300.0, 20.0, 3.0), (0.0, 0.0, 1.0), 60.0, 0.1, 'unobstructed (nu <= -0.78)'),
    (2.0, 0.01, 47.0, 0.05, 0.0, 0.0, (60.0, 30.0, 1.0), (0.0, 0.0, 1.0), 60.0, 0.1, 'LoS, nu just below -0.78'),
    (2.0, 0.01, 47.0, 0.05, 0.0, 0.0, (80.0, 30.0, 1.0), (0.0, 0.0, 1.0), 60.0, 0.1, 'LoS within Fresnel zone, kb = 0'),
    (30.0, 0.001, 47.0, 0.05, -0.1, 0.0, (35.0, 24.0, 2.0), (0.0, 0.0, 1.0), 60.0, 0.1, 'LoS within Fresnel zone, kb = 0'),
    (30.0, 0.001, 47.0, 0.05, -0.1, 0.0, (48.0, 24.0, 2.0), (0.0, 0.0, 1.0), 60.0, 0.1, 'obstructed, kb = 0'),
    (30.0, 0.001, 47.0, 0.05, -0.1, 0.0, (120.0, 24.0, 2.0), (0.0, 0.0, 1.0), 60.0, 0.1, 'obstructed, kb > 0, partial ducting'),
    (20.0, 0.01, 47.0, 0.35, 1.0, 300.0, (600.0, 20.0, 3.0), (0.0, 0.0, 1.0), 60.0, 0.1, 'obstructed, high horizon'),
    (10.0, 0.01, 20.0, 0.02, 0.0, 0.0, (120.0, 35.0, 3.0), (0.0, 0.0, 1.0), 80.0, 0.2, 'obstructed'),
    (12.0, 0.01, 47.0, 1.0, 8.0, 900.0, (1500.0, 8.0, 1.5), (0.0, 0.0, 1.0), 30.0, 0.05, 'theta > 5 deg'),
    (30.0, 0.001, 47.0, 0.05, -0.1, 0.0, (70.0, 10.0, 1.5), (110.0, 24.0, 2.0), 60.0, 0.1, 'two ridges, apex between them'),
    (20.0, 0.01, 47.0, 0.35, 1.0, 300.0, (500.0, 20.0, 3.0), (450.0, 30.0, 2.0), 60.0, 0.1, 'two ridges'),
    (6.0, 0.001, 60.0, 0.02, 0.2, 0.0, (150.0, 20.0, 2.0), (250.0, 45.0, 4.0), 60.0, 0.1, 'two ridges'),
    (40.0, 0.01, 47.0, 0.2, 3.0, 200.0, (400.0, 5.0, 0.5), (350.0, 7.0, 0.3), 60.0, 0.1, 'two sharp ridges near the station'),
    (28.0, 20.0, 47.0, 0.05, -0.1, 0.0, (120.0, 24.0, 2.0), (0.0, 0.0, 1.0), 60.0, 0.1, 'p >= beta: Ldtb = Ld'),
]
rows = []
for f, p, lat, ht, th, h0, r1, r2, dmax, st, reg in cases_s:
    d, h = R.ridges_profile(h0, [r1, r2], dmax, st)
    o = R.diffraction_ducting_slant(f, p, lat, ht, th, d, h, AE)
    rows.append((f, p, lat, ht, th, AE, h0, *r1, *r2, dmax, st, o['nu'], o['h'], o['d1'], o['J'], o['Le'], o['kb'],
                 o['Ld'], o['th_hoz'], o['d_hoz'], o['Ldtb'], 'P.526-16 §4.6; P.619-6 eq. (13); ' + reg))
write('22_terrain_slant_p526.csv', ['id', 'f_GHz', 'p_pct', 'lat_deg', 'Ht_km', 'theta_deg', 'ae_km', 'h0_m',
                                    'A1_m', 'dc1_km', 'w1_km', 'A2_m', 'dc2_km', 'w2_km', 'dmax_km', 'step_km',
                                    'nu', 'h_m', 'd1_km', 'J_dB', 'Le_dB', 'kb', 'Ld_dB', 'theta_hoz_mrad',
                                    'd_hoz_km', 'Ldtb_dB', 'source'], rows)

# 17 End-to-end §3.1 eq. (14) ---------------------------------------------
# Antennas pointing along the path; MAGRA atmosphere; p2 = 50 % (As ~ 0);
# Axp = 3 dB (§2.2 default); no terrain (Ldtb = 0).


def neg_elevation_case(lat_t, lon_t, ht, target):
    """GSO longitude giving free-space elevation `target` (deg)."""
    g = lambda lo: R.straight_line_geometry(35786.0, ht, 0.0, lat_t, wrap180(lo - lon_t))[1] - target
    return brentq(g, lon_t + 30.0, lon_t + 100.0)


scen = [  # name, f, He, Hs, lat_s, lon_s, lat_e, lon_e, e2s
    ('Bern-GSO0E', 20.0, 0.55, 35786.0, 0.0, 0.0, 46.95, 7.44),
    ('London-GSO64E', 6.425, 0.02, 35786.0, 0.0, 64.0, 51.5072, -0.1276),
    ('Tromso-GSO20W', 20.0, 0.1, 35786.0, 0.0, -20.0, 69.65, 18.96),
    ('Rio-GSO100W', 14.25, 0.0, 35786.0, 0.0, -100.0, -22.9, -43.23),
    ('Bern-LEO-high', 28.0, 0.55, 550.0, 48.0, 10.0, 46.95, 7.44),
    ('Bern-LEO-mid', 28.0, 0.55, 550.0, 52.0, 18.0, 46.95, 7.44),
    ('Bern-HAPS', 40.0, 0.55, 20.0, 47.5, 8.0, 46.95, 7.44),
    ('Jungfraujoch-GSO', 30.0, 3.5, 35786.0, 0.0, -40.0, 46.55, 7.98),
]
lon_neg = neg_elevation_case(46.5, 8.0, 2.0, -0.8)
scen.append(('Mountain-negative-elev', 20.0, 2.0, 35786.0, 0.0, lon_neg, 46.5, 8.0))

rows = []
axp = 3.0
for name, f, he, hs, la_s, lo_s, la_e, lo_e in scen:
    delta = wrap180(lo_s - lo_e)
    d, th0, _ = R.straight_line_geometry(hs, he, la_s, la_e, delta)
    th = R.fs_to_apparent(th0, he) if (he <= 3.0 and -1.0 <= th0 <= 10.0) else th0
    if th >= 0:
        ag = R.ag_ascending(f, he, hs, th)
    else:
        _, ag = R.ag_negative_elevation(f, he, hs, th)
    abs_ = R.beam_spreading_loss(th0, he) if (th0 <= 10.0 and he <= 5.0) else 0.0
    lb = R.lbfs(f, d) + axp + ag + abs_
    directions = ['E2s', 's2E'] if th >= 0 else ['E2s']
    for dirn in directions:
        rows.append((name, dirn, f, he, hs, la_s, lo_s, la_e, lo_e, 50.0, axp, d, th0, th,
                     R.lbfs(f, d), ag, abs_, 0.0, 0.0, lb))
with open(os.path.join(OUT, '17_end_to_end_eq14.csv'), 'w', newline='') as fh:
    w = csv.writer(fh, lineterminator='\n')
    w.writerow(['id', 'scenario', 'direction', 'f_GHz', 'He_km', 'Hs_km', 'lat_s_deg', 'lon_s_deg', 'lat_e_deg',
                'lon_e_deg', 'p2_pct', 'Axp_dB', 'Dts_km', 'theta0_deg', 'theta_deg', 'Lbfs_dB', 'Ag_dB', 'Abs_dB',
                'As_dB', 'Ldtb_dB', 'Lb_dB'])
    for k, r in enumerate(rows, start=1):
        w.writerow([k] + [x if isinstance(x, str) else '%.15g' % x for x in r])
print('%-32s %4d rows' % ('17_end_to_end_eq14.csv', len(rows)))

# 23 §2.8 building entry loss, Rec. ITU-R P.2109-2 ---------------------------------
rows = []
for b in (1, 2):
    for f in (0.1, 1.0, 5.0, 20.0, 60.0, 100.0):
        for p, th in [(1.0, 0.0), (10.0, 30.0), (50.0, 0.0), (50.0, 60.0), (90.0, 15.0), (99.0, -20.0)]:
            rows.append((f, p, b, th, R.bel_p2109(f, p, b, th), 'P.2109-2 eqs. (1)-(10)'))
write('23_bel_p2109.csv', ['id', 'f_GHz', 'p_pct', 'building', 'theta_deg', 'Lbe_dB', 'source'], rows)

# 24 §3.2 eq. (15), multiple-entry basic transmission loss -----------------------------
# Antennas along the path, MAGRA atmosphere, no terrain (Ldtb = 0). Lc and Lbe use the
# apparent elevation angle theta. pLc / pLbe are given (they are random in a simulation).
scen15 = [  # name, f, He, Hs, lat_s, lon_s, lat_e, lon_e
    ('Bern-GSO0E', 20.0, 0.55, 35786.0, 0.0, 0.0, 46.95, 7.44),
    ('Tromso-GSO20W', 20.0, 0.1, 35786.0, 0.0, -20.0, 69.65, 18.96),
    ('Rio-GSO100W', 14.25, 0.0, 35786.0, 0.0, -100.0, -22.9, -43.23),
    ('Bern-LEO-high', 28.0, 0.55, 550.0, 48.0, 10.0, 46.95, 7.44),
    ('Bern-LEO-mid', 28.0, 0.55, 550.0, 52.0, 18.0, 46.95, 7.44),
    ('Bern-HAPS', 40.0, 0.55, 20.0, 47.5, 8.0, 46.95, 7.44),
]
NAN = float('nan')
variants = [  # Axp, clutter, model (2: draft P.2108-2 §3.3), hs, hm, pLc, indoor, building, pLbe
    (0.0, 0, 0, NAN, NAN, NAN, 0, 0, NAN),
    (0.0, 0, 0, NAN, NAN, NAN, 1, 1, 50.0),
    (3.0, 1, 2, 5.0, 8.0, 50.0, 1, 2, 90.0),
    (0.0, 1, 2, 1.5, 5.0, 20.0, 0, 0, NAN),
    (3.0, 1, 2, 10.0, 15.0, 70.0, 0, 0, NAN),
    (3.0, 1, 2, 2.6, 15.0, 95.0, 1, 1, 30.0),
    (0.0, 1, 2, 30.0, 25.0, 50.0, 0, 0, NAN),
]
rows = []
for name, f, he, hs, la_s, lo_s, la_e, lo_e in scen15:
    delta = wrap180(lo_s - lo_e)
    d, th0, _ = R.straight_line_geometry(hs, he, la_s, la_e, delta)
    th = R.fs_to_apparent(th0, he) if (he <= 3.0 and -1.0 <= th0 <= 10.0) else th0
    ag = R.ag_ascending(f, he, hs, th)
    abs_ = R.beam_spreading_loss(th0, he) if (th0 <= 10.0 and he <= 5.0) else 0.0
    lbfs = R.lbfs(f, d)
    for axp, clut, model, hs_, hm_, plc, indoor, bld, plbe in variants:
        lc = R.clutter_p2108_2(f, th, plc, hs_, hm_) if clut else 0.0
        lbe = R.bel_p2109(f, plbe, bld, th) if indoor else 0.0
        lb = lbfs + axp + ag + abs_ + lc + lbe + 0.0
        rows.append((name, f, he, hs, la_s, lo_s, la_e, lo_e, 50.0, axp, clut, hs_, hm_, plc, indoor, bld, plbe,
                     d, th0, th, lbfs, ag, abs_, lc, lbe, 0.0, lb))
with open(os.path.join(OUT, '24_multiple_entry_eq15.csv'), 'w', newline='') as fh:
    w = csv.writer(fh, lineterminator='\n')
    w.writerow(['id', 'scenario', 'f_GHz', 'He_km', 'Hs_km', 'lat_s_deg', 'lon_s_deg', 'lat_e_deg', 'lon_e_deg',
                'p_pct', 'Axp_dB', 'clutter', 'hs_m', 'hm_m', 'pLc_pct', 'indoor', 'building', 'pLbe_pct', 'Dts_km',
                'theta0_deg', 'theta_deg', 'Lbfs_dB', 'Ag_dB', 'Abs_dB', 'Lc_dB', 'Lbe_dB', 'Ldtb_dB', 'Lb_dB'])
    for k, r in enumerate(rows, start=1):
        w.writerow([k] + [x if isinstance(x, str) else '%.15g' % x for x in r])
print('%-32s %4d rows' % ('24_multiple_entry_eq15.csv', len(rows)))

# 25 §2.7 clutter loss, draft revision P.2108-2 §3.3 (Document 3/41(Rev.2)) ------------------
# Rows sampled from the SG3 Clutter and BEL workbook values distributed with
# github.com/eeveetza/p2108 (branch dev), all rows at the boundary p = p_FcLoS, and k -> 0 cases.
P2108_CSV = os.environ.get('P2108_CSV', '/tmp/p2108/matlab/validation_examples/validation_example_p2108_3.csv')
if os.path.exists(P2108_CSV):
    wb_rows = np.loadtxt(P2108_CSV, delimiter=',', skiprows=1)
    rng = np.random.default_rng(2108)
    pick = set(rng.choice(len(wb_rows), 400, replace=False).tolist())
    for i, r in enumerate(wb_rows):              # boundary p = p_FcLoS (eq. (15) applies)
        f, th, p, hh, hm = r[:5]
        sc = 'low' if hm <= 8 else ('mid' if hm <= 20 else 'high')
        ak, bk, ck, ac, bc, av, bv = R.P2108_2_T7[sc]
        plos = R._p_state(((hh + ak) / bk) ** ck, th, ac * hh + bc, min(av * hh + bv, 100.0))
        akp, bkp, acp, b1, b2, b3, b4, ccp, avp, bvp = R.P2108_2_T8[sc]
        pfc = plos * R._p_state(akp * np.exp(bkp * hh), th, (f * 1e9) ** acp + b1 * np.exp(b2 * hh)
                                + b3 * np.exp(b4 * hh) + ccp, min(avp * hh + bvp, 0.0) * f ** -0.55 + 100.0) / 100.0
        if abs(p - pfc) < 1e-9:
            pick.add(i)
    rows = []
    for i in sorted(pick):
        f, th, p, hh, hm, lref = wb_rows[i]
        rows.append((f, th, p, hh, hm, lref, 'SG3 Clutter and BEL workbook (row %d)' % (i + 1)))
    for f, th, p, hh, hm in [(1.0, 10.0, 50.0, 2.6, 15.0), (10.0, 45.0, 90.0, 2.6, 12.0), (28.0, 70.0, 5.0, 2.6, 20.0)]:
        rows.append((f, th, p, hh, hm, R.clutter_p2108_2(f, th, p, hh, hm), 'reference, k -> 0 limit (mid-rise, h = 2.6 m)'))
    write('25_clutter_p2108_2_draft.csv', ['id', 'f_GHz', 'theta_deg', 'p_pct', 'h_m', 'hm_m', 'Lc_dB', 'source'], rows)
else:
    print('25_clutter_p2108_2_draft.csv      skipped (set P2108_CSV to the workbook values)')

# 26 Rec. ITU-R P.836-6 surface water-vapour density (annual) ----------------------------
# ITU-R SG3 validation examples (CG-3M3J-13-ValEx) and reference values computed with the
# map arrays distributed with ITU-Rpy (the maps are not part of this repository).
maps836 = R.p836_maps_from_itur()
rows = []
for r in wb['P.836-6 WV'].iter_rows(min_row=25, max_row=56, values_only=True):
    v = [x for x in r if x is not None]
    rows.append((v[0], v[1], v[2] / 1000.0, v[3], v[4], 'ITU ValEx P.836-6 WV'))
for lat, lon, alt, p in [(46.95, 7.44, 0.55, 50.0), (46.95, 7.44, 0.55, 1.0), (46.95, 7.44, 0.55, 99.0),
                         (46.55, 7.98, 3.5, 50.0), (69.65, 18.96, 0.1, 10.0), (69.65, 18.96, 0.1, 95.0),
                         (-22.9, -43.23, 0.0, 0.25), (-22.9, -43.23, 0.0, 75.0), (1.35, 103.82, 0.02, 0.1),
                         (1.35, 103.82, 0.02, 99.0), (-33.9, 151.2, 0.05, 30.0), (64.1, -21.9, 0.03, 2.5),
                         (-77.85, 166.67, 0.02, 50.0), (84.0, 20.0, 0.0, 50.0), (-84.0, 120.0, 2.8, 5.0),
                         (35.68, 139.69, 0.04, 4.0), (19.43, -99.13, 2.24, 60.0), (30.04, 31.24, 0.02, 85.0),
                         (40.71, -74.0, 0.01, 0.35), (-1.29, 36.82, 1.66, 20.0), (55.75, 37.62, 0.15, 99.0),
                         (27.99, 86.93, 5.0, 50.0), (0.0, 0.0, 0.0, 50.0), (0.0, 359.9, 0.0, 7.0),
                         (12.3, -179.9, 0.0, 0.7), (45.0, 180.0, 0.0, 15.0), (-45.0, -180.0, 0.0, 15.0),
                         (60.0, 10.0, 1.0, 97.0), (-15.8, -47.9, 1.17, 40.0), (51.5, -0.14, 0.031383, 0.2)]:
    rows.append((lat, lon, alt, p, R.p836_rho(maps836, lat, lon, p, alt), 'reference (P.836-6 Annex 1, P.1144)'))
write('26_p836_rho.csv', ['id', 'lat_deg', 'lon_deg', 'alt_km', 'p_pct', 'rho_gm3', 'source'], rows)

# 27 Ag(p1): P.835 profile scaled to rho0 (given, or from P.836 for 100 - p1 %) ----------------
scen27 = [  # name, f, He, Hs, lat_s, lon_s, lat_e, lon_e
    ('Bern-GSO0E', 20.0, 0.55, 35786.0, 0.0, 0.0, 46.95, 7.44),
    ('Tromso-GSO20W', 20.0, 0.1, 35786.0, 0.0, -20.0, 69.65, 18.96),
    ('Rio-GSO100W', 14.25, 0.0, 35786.0, 0.0, -100.0, -22.9, -43.23),
    ('Bern-LEO-mid', 28.0, 0.55, 550.0, 52.0, 18.0, 46.95, 7.44),
    ('Bern-HAPS', 40.0, 0.55, 20.0, 47.5, 8.0, 46.95, 7.44),
]
rows = []
for name, f, he, hs, la_s, lo_s, la_e, lo_e in scen27:
    delta = wrap180(lo_s - lo_e)
    d, th0, _ = R.straight_line_geometry(hs, he, la_s, la_e, delta)
    th = R.fs_to_apparent(th0, he) if (he <= 3.0 and -1.0 <= th0 <= 10.0) else th0
    for mode, p1, rho0_in in [(1, 50.0, 3.0), (1, 50.0, 15.0), (2, 0.5, NAN), (2, 5.0, NAN), (2, 50.0, NAN), (2, 95.0, NAN)]:
        if mode == 1:
            q, rho_s, rho0 = NAN, NAN, rho0_in
        else:
            q, rho_s, rho0 = R.rho0_from_p836(maps836, p1, la_e, lo_e, he)
        rows.append((name, f, he, hs, la_s, lo_s, la_e, lo_e, p1, mode, rho0_in, q, rho_s, rho0, th,
                     R.ag_ascending(f, he, hs, th, rho0)))
with open(os.path.join(OUT, '27_gas_ag_p1.csv'), 'w', newline='') as fh:
    w = csv.writer(fh, lineterminator='\n')
    w.writerow(['id', 'scenario', 'f_GHz', 'He_km', 'Hs_km', 'lat_s_deg', 'lon_s_deg', 'lat_e_deg', 'lon_e_deg',
                'p1_pct', 'mode', 'rho0_in_gm3', 'q_pct', 'rho_s_gm3', 'rho0_gm3', 'theta_deg', 'Ag_dB'])
    for k, r in enumerate(rows, start=1):
        w.writerow([k] + [x if isinstance(x, str) else '%.15g' % x for x in r])
print('%-32s %4d rows' % ('27_gas_ag_p1.csv', len(rows)))

# 28 Rec. ITU-R P.453-14 wet term of the surface refractivity Nwet (annual maps) -----------
# ITU-R SG3 validation examples (median) and reference values computed with the map arrays
# distributed with ITU-Rpy (identical to the P.453-14 files NWET_Annual_<p>.TXT).
maps453 = R.p453_maps_from_itur()
rows = []
for r in wb['P.453-14 Nwet'].iter_rows(min_row=21, max_row=40, values_only=True):
    v = [x for x in r if x is not None]
    if len(v) >= 3 and all(isinstance(x, (int, float)) for x in v[:3]):
        rows.append((v[0], v[1], 50.0, v[2], 'ITU ValEx P.453-14 Nwet'))
for lat, lon, p in [(46.95, 7.44, 50.0), (46.95, 7.44, 0.1), (46.95, 7.44, 0.7), (46.95, 7.44, 99.0),
                    (69.65, 18.96, 10.0), (69.65, 18.96, 97.0), (-22.9, -43.23, 1.0), (-22.9, 316.77, 1.0),
                    (1.35, 103.82, 0.25), (1.35, 103.82, 65.0), (-33.9, 151.2, 30.0), (35.68, 139.69, 4.0),
                    (30.04, 31.24, 85.0), (-77.85, 166.67, 50.0), (27.99, 86.93, 50.0), (90.0, 0.0, 50.0),
                    (-90.0, 45.0, 2.0), (0.0, 180.0, 50.0), (0.0, -180.0, 50.0), (12.3, 359.9, 7.0),
                    (-45.0, -179.6, 15.0), (60.0, 10.0, 0.15), (19.43, -99.13, 45.0), (55.75, 37.62, 99.0)]:
    rows.append((lat, lon, p, R.nwet_p453(maps453, lat, lon, p), 'reference (P.453-14 §2.2, P.1144 bilinear)'))
write('28_nwet_p453.csv', ['id', 'lat_deg', 'lon_deg', 'p_pct', 'Nwet', 'source'], rows)

# 29 End-to-end eq. (14) with tropospheric scintillation, Nwet = median from P.453-14 ------
# Antennas along the path; MAGRA; p1 = 50 %; Axp = 3 dB; no terrain. As from Attachment D with
# the free-space elevation theta0 (>= 5 deg) and Nwet (p = 50 %) at the earth-based station.
scen29 = [  # name, f, He, Hs, lat_s, lon_s, lat_e, lon_e
    ('Bern-GSO0E', 20.0, 0.55, 35786.0, 0.0, 0.0, 46.95, 7.44),
    ('London-GSO64E', 6.425, 0.02, 35786.0, 0.0, 64.0, 51.5072, -0.1276),
    ('Tromso-GSO20W', 20.0, 0.1, 35786.0, 0.0, -20.0, 69.65, 18.96),
    ('Rio-GSO100W', 14.25, 0.0, 35786.0, 0.0, -100.0, -22.9, -43.23),
    ('Rio-GSO100W-lon360', 14.25, 0.0, 35786.0, 0.0, 260.0, -22.9, 316.77),
    ('Singapore-GSO90E', 30.0, 0.02, 35786.0, 0.0, 90.0, 1.35, 103.82),
    ('Bern-LEO-high', 28.0, 0.55, 550.0, 48.0, 10.0, 46.95, 7.44),
]
rows = []
for name, f, he, hs, la_s, lo_s, la_e, lo_e in scen29:
    delta = wrap180(lo_s - lo_e)
    d, th0, _ = R.straight_line_geometry(hs, he, la_s, la_e, delta)
    th = R.fs_to_apparent(th0, he) if (he <= 3.0 and -1.0 <= th0 <= 10.0) else th0
    assert th0 >= 5.0
    ag = R.ag_ascending(f, he, hs, th)
    abs_ = R.beam_spreading_loss(th0, he) if (th0 <= 10.0 and he <= 5.0) else 0.0
    nwet = R.nwet_p453(maps453, la_e, lo_e, 50.0)
    for ga in (30.0, 45.0):
        for p2 in (0.01, 1.0, 10.0, 50.0, 90.0, 99.0, 99.99):
            ast = R.scintillation_ast(f, p2, th0, nwet, ga)
            rows.append((name, f, he, hs, la_s, lo_s, la_e, lo_e, p2, ga, 3.0, th0, nwet, ast,
                         R.lbfs(f, d) + 3.0 + ag + abs_ + ast))
with open(os.path.join(OUT, '29_end_to_end_scint.csv'), 'w', newline='') as fh:
    w = csv.writer(fh, lineterminator='\n')
    w.writerow(['id', 'scenario', 'f_GHz', 'He_km', 'Hs_km', 'lat_s_deg', 'lon_s_deg', 'lat_e_deg', 'lon_e_deg',
                'p2_pct', 'Ga_dBi', 'Axp_dB', 'theta0_deg', 'Nwet', 'As_dB', 'Lb_dB'])
    for k, r in enumerate(rows, start=1):
        w.writerow([k] + [x if isinstance(x, str) else '%.15g' % x for x in r])
print('%-32s %4d rows' % ('29_end_to_end_scint.csv', len(rows)))
